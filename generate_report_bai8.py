from pathlib import Path
from html import escape
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'Bao_cao_SQA_StayEasy_Theo_Khung_Bai8.docx'
ASSET = ROOT / 'bai8_diagrams'
ASSET.mkdir(exist_ok=True)

BLUE = '1F4E79'
PALE = 'EAF2F8'
GRAY = 'D9E1F2'
BLACK = RGBColor(0, 0, 0)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_border(cell, color='D9D9D9', sz='6'):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), sz)
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)

def set_cell_text(cell, text, bold=False, color=BLACK, size=10.3, align=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    if align is not None:
        p.alignment = align
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    run.font.size = Pt(size)
    run.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            set_cell_border(cell)
            if r == 0:
                set_cell_shading(cell, BLUE)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.color.rgb = RGBColor(255,255,255)
            elif r % 2 == 0:
                set_cell_shading(cell, PALE)

def tbl(caption, headers, rows, widths=None):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(caption)
    r.bold = True; r.italic = True; r.font.name = 'Times New Roman'; r.font.size = Pt(12.5)
    t = d.add_table(rows=1, cols=len(headers))
    for i, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, bold=True, color=RGBColor(255,255,255), size=10.2, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=10.0)
    style_table(t)
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    d.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def P(text='', bold=False, italic=False, center=False, size=13):
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.24
    if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = bold; r.italic = italic; r.font.name = 'Times New Roman'; r.font.size = Pt(size); r.font.color.rgb = BLACK
    return p

def bullets(items):
    for item in items:
        p = d.add_paragraph(style='List Bullet')
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(item); r.font.name='Times New Roman'; r.font.size=Pt(13); r.font.color.rgb=BLACK

def H(text, level):
    p = d.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(8 if level == 1 else 5)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text); r.font.name='Times New Roman'; r.font.color.rgb=BLACK
    return p

def diagram(name, title, nodes, arrows=None, cols=4):
    arrows = arrows or [(i, i+1) for i in range(len(nodes)-1)]
    w, h = 1100, 330
    pad = 55
    box_w = 210 if cols <= 4 else 175
    box_h = 64
    gap_x = 35
    gap_y = 85
    rows = (len(nodes) + cols - 1) // cols
    svg = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{w}" height="{h}" viewBox="0 0 {w} {h}">',
           '<defs><marker id="arrow" viewBox="0 0 10 10" refX="9" refY="5" markerWidth="6" markerHeight="6" orient="auto-start-reverse"><path d="M 0 0 L 10 5 L 0 10 z" fill="#1F4E79"/></marker></defs>',
           '<rect width="100%" height="100%" fill="white"/>']
    pos = {}
    for i, node in enumerate(nodes):
        row, col = divmod(i, cols)
        x = pad + col * (box_w + gap_x)
        y = 48 + row * (box_h + gap_y)
        pos[i] = (x, y)
        svg.append(f'<rect x="{x}" y="{y}" width="{box_w}" height="{box_h}" rx="10" fill="#EAF2F8" stroke="#1F4E79" stroke-width="2"/>')
        lines = node.split('\n')
        for j, line in enumerate(lines):
            yy = y + 26 + j * 18
            svg.append(f'<text x="{x+box_w/2}" y="{yy}" text-anchor="middle" font-family="Arial" font-size="15" fill="#000000">{escape(line)}</text>')
    for a, b in arrows:
        x1,y1=pos[a]; x2,y2=pos[b]
        if y1 == y2:
            svg.append(f'<line x1="{x1+box_w}" y1="{y1+box_h/2}" x2="{x2-6}" y2="{y2+box_h/2}" stroke="#1F4E79" stroke-width="2" marker-end="url(#arrow)"/>')
        else:
            svg.append(f'<path d="M {x1+box_w/2} {y1+box_h} L {x1+box_w/2} {y1+box_h+25} L {x2+box_w/2} {y2-7}" fill="none" stroke="#1F4E79" stroke-width="2" marker-end="url(#arrow)"/>')
    svg.append(f'<text x="{w/2}" y="315" text-anchor="middle" font-family="Arial" font-size="14" fill="#555">{escape(title)}</text></svg>')
    path = ASSET / f'{name}.svg'
    path.write_text(''.join(svg), encoding='utf-8')
    png_path = ASSET / f'{name}.png'
    canvas = Image.new('RGB', (w, h), 'white')
    draw = ImageDraw.Draw(canvas)
    try:
        font = ImageFont.truetype(r'C:\Windows\Fonts\arial.ttf', 15)
        small = ImageFont.truetype(r'C:\Windows\Fonts\arial.ttf', 14)
    except Exception:
        font = ImageFont.load_default(); small = font
    for i, node in enumerate(nodes):
        x, y = pos[i]
        draw.rounded_rectangle((x, y, x+box_w, y+box_h), radius=10, fill='#EAF2F8', outline='#1F4E79', width=2)
        lines = node.split('\n')
        bbox = draw.multiline_textbbox((0,0), '\n'.join(lines), font=font, spacing=2, align='center')
        tw = bbox[2]-bbox[0]; th = bbox[3]-bbox[1]
        draw.multiline_text((x+(box_w-tw)/2, y+(box_h-th)/2-2), '\n'.join(lines), font=font, fill='black', spacing=2, align='center')
    for a, b in arrows:
        x1,y1=pos[a]; x2,y2=pos[b]
        if y1 == y2:
            start=(x1+box_w,y1+box_h//2); end=(x2-5,y2+box_h//2)
            draw.line((start[0],start[1],end[0],end[1]), fill='#1F4E79', width=3)
            draw.polygon([(end[0],end[1]),(end[0]-10,end[1]-6),(end[0]-10,end[1]+6)], fill='#1F4E79')
        else:
            sx=x1+box_w//2; sy=y1+box_h; ex=x2+box_w//2; ey=y2-5
            mid=sy+22
            draw.line((sx,sy,sx,mid,ex,ey), fill='#1F4E79', width=3)
            draw.polygon([(ex,ey),(ex-6,ey-10),(ex+6,ey-10)], fill='#1F4E79')
    draw.text((w//2, 305), title, font=small, fill='#555555', anchor='ma')
    canvas.save(png_path, 'PNG')
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(1)
    run = p.add_run(); run.add_picture(str(png_path), width=Inches(6.7))
    cp = d.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER; cp.paragraph_format.space_after=Pt(5)
    rr = cp.add_run(title); rr.italic=True; rr.font.name='Times New Roman'; rr.font.size=Pt(11.5); rr.font.color.rgb=BLACK

def page_break():
    d.add_page_break()

d = Document()
sec = d.sections[0]
sec.top_margin = Inches(0.80); sec.bottom_margin = Inches(0.75); sec.left_margin = Inches(0.90); sec.right_margin = Inches(0.75)
styles = d.styles
styles['Normal'].font.name='Times New Roman'; styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman'); styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman'); styles['Normal'].font.size=Pt(13); styles['Normal'].font.color.rgb=BLACK
for level in (1,2,3,4):
    st = styles[f'Heading {level}']; st.font.name='Times New Roman'; st._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman'); st._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman'); st.font.color.rgb=BLACK; st.font.bold=True; st.font.size=Pt(15 if level==1 else 14 if level==2 else 13)
    st.paragraph_format.keep_with_next=True
styles['Title'].font.name='Times New Roman'; styles['Title']._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman'); styles['Title']._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman'); styles['Title'].font.color.rgb=BLACK

# header/footer
header = sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.CENTER
rh=header.add_run('STAYEASY HOMESTAY  |  BÁO CÁO ĐẢM BẢO CHẤT LƯỢNG PHẦN MỀM'); rh.font.name='Times New Roman'; rh.font.size=Pt(8.5); rh.font.color.rgb=RGBColor(100,100,100)
footer = sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
rf=footer.add_run('Trang '); rf.font.name='Times New Roman'; rf.font.size=Pt(9)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)

# Cover
P('TRƯỜNG ĐẠI HỌC ................................................', center=True, size=13)
P('KHOA / BỘ MÔN ................................................', center=True, size=13)
P('', center=True)
p=d.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(12); r=p.add_run('BÁO CÁO ĐẢM BẢO CHẤT LƯỢNG PHẦN MỀM'); r.font.name='Times New Roman'; r.font.size=Pt(19); r.font.bold=True; r.font.color.rgb=BLACK
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('HỆ THỐNG ĐẶT PHÒNG HOMESTAY STAYEASY'); r.font.name='Times New Roman'; r.font.size=Pt(16); r.bold=True; r.font.color.rgb=BLACK
P('Áp dụng mô hình Waterfall và tiêu chuẩn ISO/IEC 25010:2011', center=True, italic=True, size=13)
P('\nSinh viên thực hiện: ........................................................\nLớp: ........................................................\nGiảng viên hướng dẫn: ........................................................', center=True, size=13)
P('\nTP. Hồ Chí Minh, tháng .... năm ....', center=True, size=13)
page_break()

# TOC
H('MỤC LỤC', 1)
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Mục lục được cập nhật tự động trong Microsoft Word.'); r.font.name='Times New Roman'; r.font.size=Pt(11); r.italic=True
p=d.add_paragraph(); fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'TOC \\o "1-4" \\h \\z \\u'); p._p.append(fld)
page_break()

# 1
H('1. GIỚI THIỆU VỀ DỰ ÁN PHẦN MỀM',1)
H('1.1. Tổng quan về dự án',2)
P('StayEasy Homestay là hệ thống web hỗ trợ quy trình tìm kiếm, đặt phòng và quản lý hoạt động lưu trú tại homestay. Hệ thống hướng tới hai nhóm người dùng chính là khách hàng và quản trị viên; mỗi nhóm được cung cấp giao diện và quyền truy cập phù hợp với trách nhiệm nghiệp vụ. Mục tiêu của dự án là tập trung hóa dữ liệu phòng, đơn đặt, hồ sơ lưu trú và lịch sử hoạt động trong một nền tảng có thể truy xuất, kiểm soát và bảo trì.')
P('Hệ thống được triển khai theo kiến trúc web nhiều tầng. Thành phần thứ nhất là cổng khách hàng, gồm các trang HTML/CSS/JavaScript phục vụ đăng ký, đăng nhập, xem loại phòng, xem căn phòng, gửi yêu cầu đặt phòng, theo dõi đơn, cập nhật hồ sơ và xem lịch sử lưu trú. Thành phần thứ hai là cổng quản trị, gồm dashboard và các màn hình quản lý tài khoản, loại phòng, căn phòng, đơn đặt, hồ sơ, lịch sử và nhật ký hoạt động. Thành phần thứ ba là lớp dịch vụ phía máy chủ, gồm Express trên Node.js, cơ sở dữ liệu MySQL và dịch vụ chatbot FastAPI tùy chọn để tư vấn thông tin homestay.')
P('Các quy tắc nghiệp vụ quan trọng gồm: khách hàng chỉ được xem và thay đổi dữ liệu thuộc quyền sở hữu; đơn đặt phải gắn với khách, căn phòng, khoảng thời gian và trạng thái hợp lệ; dữ liệu phòng không được tạo quan hệ sai với loại phòng; quản trị viên phải được xác thực và kiểm tra vai trò trước thao tác quản trị; mật khẩu không được lưu ở dạng văn bản; chatbot chỉ là dịch vụ hỗ trợ và không được làm gián đoạn luồng đặt phòng khi không sẵn sàng. Yêu cầu khác gồm khả năng cấu hình qua biến môi trường, truy xuất nhật ký, khôi phục cơ sở dữ liệu và có tài liệu triển khai.')
P('Về ranh giới dữ liệu, trình duyệt chỉ giao tiếp với REST API; frontend không truy cập trực tiếp MySQL. Backend chịu trách nhiệm xác thực, kiểm tra dữ liệu đầu vào, thực hiện nghiệp vụ và trả response. Database chịu trách nhiệm lưu trữ và ràng buộc quan hệ. Chatbot FastAPI là một boundary phụ: các yêu cầu tư vấn có thể được chuyển tới chatbot, nhưng các thao tác tạo đơn, hủy đơn hoặc cập nhật hồ sơ vẫn phải đi qua API chính và middleware của StayEasy.')
P('Các giả định của báo cáo là môi trường kiểm thử có MySQL theo database.sql, backend Express chạy được bằng package.json, chatbot chạy độc lập khi được bật, tài khoản admin và customer được tạo riêng, dữ liệu phòng có đủ trạng thái để kiểm thử, và các thành viên sử dụng cùng một quy ước version. Khi một giả định không đúng, QA phải ghi nhận trong Test Plan hoặc Defect Report thay vì kết luận chất lượng dựa trên môi trường không xác định.')
P('Về công nghệ, backend sử dụng Node.js/Express và mysql2 connection pool; frontend sử dụng HTML, CSS và JavaScript theo từng module; database.sql mô tả các bảng QUANTRIVIEN, KHACHHANG, LOAIPHONG, CANPHONG, DATPHONG, HOSOLUUTRU, LICHSULUUTRU và hoatdong; server.py cung cấp chatbot FastAPI tùy chọn. Cơ chế bảo vệ hiện có gồm JWT, PBKDF2-SHA256, rate limit đăng nhập, CORS whitelist, truy vấn tham số hóa, kiểm tra ownership và quyền admin.')
H('1.2. Yêu cầu người dùng và phạm vi dự án',2)
H('1.2.1. Yêu cầu người dùng',3)
P('Nhóm người dùng và nhu cầu tương ứng được xác định như sau:')
bullets(['Khách hàng: đăng ký tài khoản; đăng nhập và đăng xuất; xem loại phòng, căn phòng, tiện nghi và giá; chọn ngày/giờ; gửi yêu cầu đặt phòng; theo dõi trạng thái đơn; hủy đơn theo quy định; cập nhật hồ sơ; đổi mật khẩu; xem hồ sơ và lịch sử lưu trú; sử dụng chatbot để hỏi thông tin homestay.','Quản trị viên: đăng nhập với vai trò quản trị; quản lý loại phòng; quản lý căn phòng; quản lý tài khoản khách hàng; xem và cập nhật đơn đặt; quản lý hồ sơ lưu trú; quản lý lịch sử lưu trú; xem dashboard; xem nhật ký hoạt động; kiểm tra dữ liệu trước khi xóa hoặc cập nhật.','Nhóm phát triển và QA: truy xuất yêu cầu tới API, database và test case; review tài liệu, thiết kế và code; quản lý defect; kiểm soát baseline, thay đổi, kiểm thử và hồ sơ nghiệm thu.'])
H('1.2.2. Phạm vi dự án',3)
H('1.2.2.1. Phạm vi chức năng',4)
P('Phạm vi chức năng bao gồm xác thực tài khoản; tra cứu loại phòng và căn phòng; đặt phòng theo khoảng thời gian; kiểm tra tình trạng đã đặt; xem, hủy và theo dõi đơn thuộc khách hàng; cập nhật hồ sơ và đổi mật khẩu; các chức năng CRUD quản trị; dashboard và nhật ký hoạt động; chatbot tư vấn thông tin tĩnh hoặc thông tin được cấu hình cho homestay.')
H('1.2.2.2. Phạm vi phi chức năng',4)
P('Tính khả dụng: người dùng có thể hoàn thành các luồng chính với thông báo lỗi dễ hiểu. Hiệu năng: các API trọng yếu có thời gian phản hồi được đo trong Test Plan, ngưỡng chấp nhận để trống cho kết quả thực tế. Bảo mật: JWT, PBKDF2-SHA256, rate limit, CORS, parameterized query, kiểm tra role/ownership, chống truy cập IDOR và escape dữ liệu giao diện. Tương thích: kiểm tra trên Chrome, Edge, Firefox và các kích thước viewport được thống nhất trong Test Plan. Giao diện: responsive, biểu mẫu có validation, thông báo trạng thái rõ ràng. CSDL: khóa chính, khóa ngoại, ràng buộc duy nhất, trạng thái nghiệp vụ và tính toàn vẹn dữ liệu phải được kiểm tra.')
H('1.2.2.3. Phạm vi không bao gồm',4)
P('Phạm vi hiện tại không bao gồm thanh toán trực tuyến, kết nối OTA hoặc nhiều chi nhánh, SSO/MFA, hạ tầng production đa vùng, hệ thống WAF chuyên dụng, tự động hóa hoàn toàn quy trình bằng hội thoại và các cam kết SLA chưa được phê duyệt trong yêu cầu dự án.')
H('1.3. Yêu cầu phần mềm',2)
P('Các yêu cầu được phân loại thành yêu cầu nghiệp vụ BR, yêu cầu hệ thống SR và yêu cầu phi chức năng NFR. Mỗi yêu cầu có mã duy nhất, có thể truy xuất tới module, tiêu chí chấp nhận và test case.')
tbl('Bảng 1. Yêu cầu nghiệp vụ (BR)', ['Mã yêu cầu','Nội dung yêu cầu'], [['BR01','Khách hàng tạo tài khoản bằng thông tin hợp lệ và đăng nhập được.'],['BR02','Khách hàng tra cứu loại phòng, căn phòng, tiện nghi và giá.'],['BR03','Khách hàng chọn thời gian và gửi yêu cầu đặt phòng.'],['BR04','Khách hàng xem và hủy đơn thuộc quyền sở hữu theo trạng thái được phép.'],['BR05','Khách hàng cập nhật hồ sơ và đổi mật khẩu.'],['BR06','Quản trị viên quản lý loại phòng và căn phòng.'],['BR07','Quản trị viên quản lý tài khoản, đơn đặt, hồ sơ và lịch sử lưu trú.'],['BR08','Quản trị viên xem dashboard và nhật ký hoạt động.'],['BR09','Khách hàng sử dụng chatbot để nhận tư vấn thông tin homestay.']],[0.9,5.9])
tbl('Bảng 2. Yêu cầu hệ thống (SR)', ['Mã yêu cầu','Nội dung yêu cầu'], [['SR01','Backend cung cấp REST API có mã HTTP, dữ liệu trả về và thông báo lỗi rõ ràng.'],['SR02','Middleware xác thực token trước route bảo vệ và kiểm tra role admin khi cần.'],['SR03','API kiểm tra ownership trước khi khách hàng xem, sửa hoặc hủy dữ liệu.'],['SR04','MySQL duy trì khóa chính, khóa ngoại, ràng buộc duy nhất và trạng thái dữ liệu.'],['SR05','Dữ liệu đầu vào được validation trước khi thực hiện nghiệp vụ hoặc truy vấn.'],['SR06','Cấu hình port, database và JWT secret được đọc từ biến môi trường.'],['SR07','Hệ thống có nhật ký hoạt động cho các thao tác quản trị quan trọng.'],['SR08','Chatbot được gọi qua API riêng và có xử lý timeout/fallback.']],[0.9,5.9])
tbl('Bảng 3. Yêu cầu phi chức năng (NFR)', ['Mã yêu cầu','Nội dung yêu cầu'], [['NFR01','Hiệu năng API trọng yếu được đo theo môi trường và dữ liệu quy định trong Test Plan.'],['NFR02','Không có tài khoản khách hàng truy cập được route admin hoặc dữ liệu khách hàng khác.'],['NFR03','Giao diện hiển thị và thao tác được trên Chrome, Edge, Firefox và viewport mobile.'],['NFR04','Mật khẩu được bảo vệ bằng PBKDF2-SHA256 với salt; secret không lưu trong repository.'],['NFR05','Truy vấn dữ liệu sử dụng parameterized query và xử lý lỗi an toàn.'],['NFR06','Mã nguồn có cấu trúc module, có review và tài liệu README phục vụ bảo trì.'],['NFR07','Có thể cài đặt lại bằng package.json, requirements.txt và database.sql.'],['NFR08','Kết quả kiểm thử, defect và quyết định release được lưu thành bằng chứng.']],[0.9,5.9])
H('1.4. Đội dự án và đảm bảo chất lượng',2)
P('Nhân sự và trách nhiệm chất lượng được tổ chức theo mô hình phân quyền của báo cáo SQA.')
tbl('Bảng 4. Thành viên đội dự án', ['TT','Vai trò','Họ tên','Ghi chú'], [['1','PM','...','Quản lý phạm vi, kế hoạch, baseline và phê duyệt cổng chất lượng'],['2','QA leader','...','Lập kế hoạch QA, điều phối review, test và báo cáo'],['3','SQA team','...','Theo dõi quy trình, checklist, defect, bằng chứng và cải tiến'],['4','Developer','...','Lập trình backend/frontend, unit test và sửa defect'],['5','Trưởng nhóm thiết kế','...','Thiết kế kiến trúc, API, dữ liệu và review SDD'],['6','Khách hàng nội bộ','...','Xác nhận yêu cầu, tiêu chí chấp nhận và UAT']],[0.4,1.6,1.4,3.4])

# 2
H('2. THIẾT LẬP MỤC TIÊU VÀ YÊU CẦU CHẤT LƯỢNG CHO DỰ ÁN',1)
H('2.1. Mục tiêu chất lượng',2)
P('Hệ thống đảm bảo chất lượng của StayEasy được xây dựng dựa trên ISO/IEC 25010:2011 về mô hình chất lượng sản phẩm phần mềm; IEEE 730 về kế hoạch đảm bảo chất lượng phần mềm; ISO/IEC 12207 về các quy trình vòng đời phần mềm; và CMMI-DEV v2.0 về quản lý, đo lường, đảm bảo chất lượng và cải tiến quy trình. Các chuẩn được sử dụng làm cơ sở để chuyển yêu cầu thành tiêu chí kiểm tra, sản phẩm đầu ra, trách nhiệm và bằng chứng.')
bullets(['Đảm bảo tính phù hợp chức năng: các luồng đăng nhập, tra cứu, đặt phòng, quản trị và chatbot trong phạm vi phải được đặc tả, kiểm thử và truy xuất.','Đảm bảo an toàn dữ liệu: xác thực, phân quyền, ownership, bảo vệ mật khẩu và truy vấn dữ liệu phải được review và kiểm thử tiêu cực.','Đảm bảo độ tin cậy: lỗi đầu vào, lỗi database, timeout chatbot và lỗi dịch vụ phải được xử lý có kiểm soát, không tạo dữ liệu sai.','Đảm bảo khả dụng: giao diện, thông báo, validation và luồng thao tác phải phù hợp với khách hàng và quản trị viên.','Đảm bảo hiệu năng: thời gian phản hồi API, tải truy cập và các điểm nghẽn phải được đo bằng công cụ phù hợp, không ghi số liệu khi chưa thực thi.','Đảm bảo khả bảo trì: code, tài liệu, cấu trúc module, review record, test case và change log phải được quản lý theo baseline.','Đảm bảo tính di động và khả năng bàn giao: dự án có thể cài đặt lại bằng cấu hình, mã nguồn, database script và hướng dẫn được kiểm soát phiên bản.'])
H('2.2. Yêu cầu chất lượng',2)
P('Các yêu cầu chất lượng được xây dựng theo tám đặc tính chính của ISO/IEC 25010:2011. Với mỗi đặc tính, báo cáo xác định yêu cầu cụ thể và bằng chứng cần thu thập trong các hoạt động review, kiểm thử, nghiệm thu và bàn giao.')
P('Cách đánh giá được chia thành ba lớp. Lớp thứ nhất là đánh giá tài liệu và thiết kế, tập trung vào tính đầy đủ, nhất quán, khả kiểm thử và khả truy xuất. Lớp thứ hai là đánh giá mã nguồn và cấu hình, tập trung vào lỗi logic, lỗi xử lý ngoại lệ, rủi ro bảo mật, chất lượng cấu trúc và khả năng bảo trì. Lớp thứ ba là kiểm thử động và nghiệm thu, tập trung vào hành vi người dùng, dữ liệu sau thao tác, response API, trạng thái giao diện và các điều kiện lỗi.')
P('Mọi tiêu chí đo được phải gắn với phương pháp đo, môi trường, dữ liệu, người thực hiện và bằng chứng. Ví dụ, hiệu năng phải ghi endpoint, số lượng request, cấu hình máy, dữ liệu seed và kết quả; security phải ghi token/role, bước tái hiện và response; usability phải ghi luồng, viewport và nhận xét của người dùng. Cách ghi này giúp phân biệt mục tiêu chất lượng với kết quả đã thực thi và tránh kết luận không có cơ sở.')
tbl('Bảng 5. Yêu cầu chất lượng theo ISO/IEC 25010:2011', ['Đặc tính chính','Yêu cầu chất lượng cụ thể'], [['Functional suitability','Đăng nhập, tra cứu, đặt phòng, xem/hủy đơn, quản trị dữ liệu, dashboard và chatbot thực hiện đúng phạm vi; yêu cầu có tiêu chí chấp nhận và test case liên kết.'],['Performance efficiency','API tra cứu, đăng nhập, đặt phòng và dashboard được đo trên môi trường kiểm thử; kết quả ghi rõ dữ liệu, tải, thời gian và ngưỡng được phê duyệt.'],['Compatibility','Frontend hoạt động nhất quán trên Chrome, Edge, Firefox và viewport desktop/mobile; chatbot và backend giao tiếp đúng API contract.'],['Usability','Luồng đặt phòng có điều hướng rõ, biểu mẫu có validation, thông báo lỗi dễ hiểu và trạng thái thao tác được phản hồi.'],['Reliability','Lỗi database, token, dữ liệu thiếu, timeout chatbot và thao tác lặp được kiểm soát; không tạo bản ghi sai sau lỗi.'],['Security','JWT, role, ownership, PBKDF2-SHA256, rate limit, CORS, parameterized query, XSS/IDOR và secret môi trường được kiểm tra.'],['Maintainability','Module backend/frontend rõ trách nhiệm; code có guideline, review, test liên quan, tài liệu và change log.'],['Portability','Cấu hình được tách khỏi mã nguồn; dự án có thể khởi tạo lại bằng Node.js, Python/FastAPI, MySQL, package.json và database.sql.']],[1.7,5.1])
page_break()

# 3
H('3. XÁC ĐỊNH MÔ HÌNH PHÁT TRIỂN',1)
P('Mô hình được lựa chọn cho StayEasy là Waterfall, một mô hình phát triển tuần tự trong đó yêu cầu, đặc tả, thiết kế, lập trình, kiểm thử và bàn giao được tổ chức thành các giai đoạn có đầu ra và cổng kiểm soát. Waterfall không thuộc nhóm Agile; các thay đổi sau baseline được xử lý thông qua Change Request và đánh giá tác động.')
P('Về mức độ rõ ràng của yêu cầu, các chức năng cốt lõi của StayEasy đã được định hình tương đối rõ từ mã nguồn và database: tài khoản, loại phòng, căn phòng, đặt phòng, hồ sơ, lịch sử và quản trị. Về quy mô và độ phức tạp, hệ thống có nhiều module nhưng phạm vi đồ án được giới hạn, phù hợp với việc chốt SRS, SDD và test plan trước khi lập trình hoặc kiểm thử. Về thời gian và ngân sách, nhóm cần dự toán theo các mốc tuần tự để kiểm soát tiến độ và tránh chi phí phát sinh do thay đổi không đánh giá.')
P('Về tương tác khách hàng, khách hàng nội bộ tham gia xác nhận yêu cầu ở đầu kỳ, duyệt tiêu chí chấp nhận và thực hiện UAT ở cuối kỳ; việc trao đổi không làm mất cổng kiểm soát. Về mức độ kiểm soát quy trình, Waterfall giúp nhóm xác định rõ artifact, reviewer, điều kiện chuyển giai đoạn và bằng chứng. Đây là ưu điểm quan trọng đối với báo cáo SQA, nơi cần chứng minh liên kết từ yêu cầu đến kiểm thử.')
P('So sánh ngắn: Agile/Scrum phù hợp khi yêu cầu thay đổi liên tục nhưng cần năng lực quản trị sprint và phản hồi thường xuyên; V-model nhấn mạnh liên kết giữa từng giai đoạn phát triển và cấp độ kiểm thử, có thể dùng bổ sung trong Test Plan; Spiral phù hợp dự án có rủi ro lớn cần lặp phân tích rủi ro; Incremental chia sản phẩm thành các increment; RAD ưu tiên tốc độ tạo prototype. Với phạm vi, tài nguyên và mục tiêu kiểm soát hồ sơ của StayEasy, Waterfall là lựa chọn chính; các kỹ thuật review, test tự động và feedback vẫn có thể được dùng bên trong từng giai đoạn.')
P('Các cổng chất lượng của Waterfall được xác định như sau: sau phân tích phải có phạm vi và nhóm người dùng; sau đặc tả phải có SRS, mã yêu cầu và tiêu chí chấp nhận; sau thiết kế phải có SDD, API contract và mô hình dữ liệu; sau lập trình phải có code review và unit test; sau kiểm thử phải có defect status, regression và Test Summary; trước bàn giao phải có UAT, release package, hướng dẫn triển khai và hồ sơ khôi phục. Nếu cổng chưa đạt, artifact quay lại bước có liên quan thông qua finding hoặc Change Request.')
P('Việc chọn Waterfall không có nghĩa là bỏ qua phản hồi. Phản hồi được đặt tại các điểm kiểm soát: khách hàng nội bộ xác nhận SRS, reviewer xác nhận SDD/code, QA phản hồi kết quả test và người dùng xác nhận UAT. Điểm khác biệt là mỗi phản hồi được ghi nhận thành finding, defect hoặc CR, có owner và quyết định, thay vì thay đổi trực tiếp không có truy vết.')
page_break()

# 4
H('4. THIẾT LẬP CÁC TIÊU CHUẨN, THỦ TỤC VÀ TÀI LIỆU QA',1)
H('4.1. Tiêu chuẩn tài liệu',2)
P('Bộ tài liệu QA của StayEasy được quản lý theo mã tài liệu, phiên bản, người lập, reviewer, trạng thái và ngày thay đổi. Tài liệu chỉ được dùng làm baseline khi đã được review và phê duyệt. Danh mục gồm 11 tài liệu sau:')
P('Quan hệ giữa các tài liệu phải được duy trì nhất quán. SRS là nguồn của yêu cầu và tiêu chí chấp nhận; SDD giải thích cách hiện thực yêu cầu; Test_Plan xác định chiến lược và môi trường; Test_Cases kiểm tra yêu cầu; Defect_Report ghi sai khác; Test_Summary_Report tổng hợp quyết định; Retrospective ghi bài học. Khi một yêu cầu hoặc thiết kế thay đổi, QA phải kiểm tra tác động tới các tài liệu liên kết trước khi cập nhật baseline.')
P('Mỗi tài liệu tối thiểu có phần nhận diện gồm tên, mã, phiên bản, trạng thái, tác giả, reviewer, ngày lập và lịch sử thay đổi. Các bảng, hình và thuật ngữ phải được đánh số thống nhất. Bản đang soạn không được dùng làm bằng chứng release; chỉ bản được phê duyệt và có thể truy xuất tới version mã nguồn mới được dùng trong Test Summary hoặc UAT.')
bullets(['(1) QA_Plan: mục tiêu, tiêu chuẩn, vai trò, kế hoạch review, kiểm thử, báo cáo và tiêu chí chấp nhận.','(2) Dev_Workflow_Design: mô tả quy trình Waterfall, Hình 1, Bảng 6 và Bảng 7.','(3) Doc_Stand_Guide: quy tắc đặt tên, định dạng, phiên bản, review và lưu trữ tài liệu.','(4) SRS: yêu cầu nghiệp vụ, hệ thống, phi chức năng và tiêu chí chấp nhận.','(5) SDD: kiến trúc, module, API, dữ liệu, bảo mật và quyết định thiết kế.','(6) Test_Plan: phạm vi, chiến lược, môi trường, dữ liệu, lịch, vai trò và tiêu chí kết thúc.','(7) Test_Cases: ID, yêu cầu liên kết, tiền điều kiện, bước, dữ liệu, expected, actual và evidence.','(8) Code_Review_Checklist: tiêu chí đọc hiểu, bảo mật, xử lý lỗi, test và maintainability.','(9) Defect_Report: mã lỗi, môi trường, bước tái hiện, severity, owner, bằng chứng và trạng thái.','(10) Test_Summary_Report: tổng hợp thực thi, pass/fail, defect, rủi ro tồn đọng và quyết định.','(11) Retrospective: bài học, nguyên nhân, hành động cải tiến và người chịu trách nhiệm.'])
P('Dev_Workflow_Design của StayEasy áp dụng Waterfall theo luồng tuần tự sau:')
diagram('fig1_waterfall','Hình 1. Sơ đồ luồng công việc chính trong quy trình Waterfall', ['Phân tích\nyêu cầu','Đặc tả\nSRS/RTM','Thiết kế\nSDD/API/ERD','Lập trình\nCode/Review','Kiểm thử\nTest/Defect','Bàn giao\nUAT/Release'], cols=3)
tbl('Bảng 6. Mô tả quy trình làm việc theo từng giai đoạn',['Giai đoạn','Hoạt động chính','Tài liệu đầu ra','Người chịu trách nhiệm'], [['1. Phân tích','Thu thập nghiệp vụ, nhóm người dùng, phạm vi và ràng buộc','Danh sách yêu cầu, phạm vi dự án','PM, BA, khách hàng nội bộ'],['2. Đặc tả yêu cầu','Mã hóa BR/SR/NFR, tiêu chí chấp nhận, RTM','SRS, RTM','BA, QA leader, SQA team'],['3. Thiết kế','Thiết kế kiến trúc, API, CSDL, giao diện và bảo mật','SDD, ERD, API contract','Trưởng nhóm thiết kế, Developer'],['4. Lập trình','Phát triển module, unit test, review và quản lý branch','Mã nguồn, review record','Developer, QA leader'],['5. Kiểm thử','Thực thi unit/integration/system/security/acceptance, mở defect','Test Cases, Defect Report, Test Summary','QA/Tester, Developer, khách hàng nội bộ'],['6. Bàn giao','Đóng lỗi, xác nhận UAT, triển khai, hướng dẫn và lưu baseline','Release package, UAT sign-off, README','PM, QA leader, SQA team']],[1.0,2.4,2.0,1.4])
tbl('Bảng 7. Phân công vai trò thành viên đội dự án',['Vai trò','Nhiệm vụ trong quy trình','Sản phẩm/bằng chứng'], [['PM','Phê duyệt phạm vi, kế hoạch, CR, baseline và quyết định release','Kế hoạch, biên bản, sign-off'],['QA leader','Điều phối QA, review, test plan, defect triage và test summary','QA_Plan, Test Plan, report'],['SQA team','Audit việc tuân thủ, kiểm tra checklist, lưu bằng chứng và cải tiến','Checklist, audit record, retrospective'],['Developer','Thiết kế/lập trình, unit test, self-check, sửa defect','Code, unit result, commit, fix response'],['Trưởng nhóm thiết kế','Chủ trì SDD, ERD, API contract và design review','SDD, ERD, review record'],['Khách hàng nội bộ','Xác nhận yêu cầu, tiêu chí chấp nhận và UAT','Requirement sign-off, UAT record']],[1.3,3.5,2.0])
H('4.2. Thiết lập chuẩn mã nguồn',2)
P('(1) Quy tắc chung: backend sử dụng JavaScript chạy trên Node.js/Express; chatbot sử dụng Python/FastAPI; database sử dụng MySQL; frontend sử dụng HTML/CSS/JavaScript. Phiên bản runtime, package và database phải được ghi trong README hoặc Test Plan. IDE được nhóm thống nhất, không commit thư mục build, secret hoặc dữ liệu cá nhân vào repository.')
P('(2) Quy tắc đặt tên: biến và hàm JavaScript dùng camelCase; lớp hoặc constructor dùng PascalCase; hằng số dùng UPPER_SNAKE_CASE; route và tên file module phản ánh miền nghiệp vụ; tên bảng và cột nhất quán với database.sql; Python dùng snake_case cho hàm/biến và PascalCase cho class. Tên phải mô tả mục đích, tránh viết tắt khó hiểu.')
P('(3) Quy tắc mã nguồn: module có trách nhiệm đơn; hàm nghiệp vụ nên ngắn và dễ test; logic lặp phải được tách helper; complexity cao phải được review; comment dùng để giải thích lý do hoặc quy tắc nghiệp vụ, không lặp lại câu lệnh; mọi I/O, database và gọi dịch vụ ngoài phải có try-catch hoặc cơ chế xử lý lỗi tương đương; input phải validation; truy vấn phải parameterized; log không chứa mật khẩu/token; thay đổi chính phải có test và review.')
P('Đối với backend StayEasy, review cần tập trung vào các vùng rủi ro cao: middleware authenticateToken và requireAdmin; các route có tham số idKH, id căn hoặc id đơn; logic kiểm tra trùng ngày; thao tác cập nhật/xóa; xử lý lỗi connection pool; và các endpoint ghi nhật ký. Đối với frontend, review tập trung vào việc lấy token, đồng bộ trạng thái đăng nhập, hiển thị response lỗi, escape dữ liệu và không để lộ quyền admin trong giao diện. Đối với chatbot, review tập trung vào timeout, dữ liệu trả về, xử lý service unavailable và giới hạn phạm vi tư vấn.')
P('Quy ước Git gồm branch cho thay đổi, commit mô tả một mục tiêu, pull request có liên kết yêu cầu/defect, reviewer độc lập, tag cho baseline và release note cho bản bàn giao. Không sửa trực tiếp branch baseline nếu chưa có phê duyệt. Khi cần hotfix, nhóm phải ghi lý do, phạm vi, test hồi quy và cập nhật tài liệu sau khi hệ thống ổn định.')
P('Tài liệu tham khảo được áp dụng gồm Google JavaScript Style Guide cho JavaScript, PEP 8 cho Python và Git Best Practices cho branch, commit, pull request, tag và release. Các hướng dẫn được điều chỉnh theo cấu trúc StayEasy và phải được kiểm tra trong code review.')
H('4.3. Thiết lập quy trình review',2)
P('Quy trình review gồm 6 bước: (1) xác định artifact, version, phạm vi và reviewer; (2) tác giả tự kiểm tra bằng checklist; (3) reviewer độc lập đọc và ghi finding; (4) phân loại finding theo blocker, major, minor hoặc observation; (5) tác giả sửa, phản hồi và cung cấp diff/evidence; (6) reviewer hoặc QA xác nhận, đóng finding và cho phép baseline.')
P('Lưu ý: peer review áp dụng cho SRS, SDD, Test Plan, Test Cases, Defect Report và các tài liệu quan trọng; code review bắt buộc khi thay đổi lớn, đặc biệt khi thay đổi trên 30% mã của module, ảnh hưởng bảo mật, database hoặc API contract; reviewer phải độc lập với người viết thay đổi; finding chưa có bằng chứng xử lý không được xem là đã đóng.')
P('Mức độ finding được diễn giải thống nhất: blocker ngăn tiếp tục hoặc phát hành; major làm sai nghiệp vụ, bảo mật hoặc dữ liệu quan trọng; minor ảnh hưởng cục bộ nhưng có workaround; observation là đề xuất cải tiến chưa tạo lỗi trực tiếp. Review record phải ghi artifact/version, phạm vi, danh sách reviewer, finding, severity, owner, trạng thái và liên kết commit hoặc tài liệu phản hồi.')
P('Review không chỉ là đọc code. Với SRS, reviewer kiểm tra tính đầy đủ và không mơ hồ; với SDD, kiểm tra tính nhất quán giữa API, ERD và module; với Test Cases, kiểm tra coverage, expected và dữ liệu biên; với Defect Report, kiểm tra khả năng tái hiện và evidence. QA leader dùng các review record này để xác nhận cổng chất lượng trước khi chuyển giai đoạn.')
H('4.4. Thiết lập quy trình kiểm thử',2)
P('Quy trình kiểm thử gồm 7 bước: (1) phân tích yêu cầu và rủi ro; (2) lập chiến lược, phạm vi và môi trường; (3) thiết kế test case, dữ liệu và expected result; (4) chuẩn bị build, database seed, tài khoản và trình duyệt; (5) thực thi và ghi actual result/evidence; (6) mở defect, re-test và regression; (7) tổng hợp Test Summary, rủi ro tồn đọng và quyết định acceptance/release.')
P('Các cấp độ kiểm thử gồm unit test cho helper, validation và logic nhỏ; integration test cho frontend/API/MySQL; system test cho luồng nghiệp vụ hoàn chỉnh; security test cho token, role, ownership, SQLi, XSS và rate limit; compatibility test cho trình duyệt/viewport; acceptance test do khách hàng nội bộ xác nhận. Acceptance có thể tổ chức theo alpha trong môi trường nhóm và beta khi có nhóm người dùng thử nghiệm được phê duyệt.')
P('Dữ liệu kiểm thử phải tách khỏi dữ liệu thật và có trạng thái đủ để kiểm tra các nhánh nghiệp vụ. Bộ dữ liệu tối thiểu gồm customer hợp lệ, admin hợp lệ, token sai/hết hạn, số điện thoại trùng, mật khẩu sai, căn phòng còn trống, căn phòng đã có đơn, ngày thiếu, ngày không hợp lệ, id không tồn tại và dữ liệu có ký tự đặc biệt. Sau mỗi lần test, dữ liệu được reset hoặc ghi rõ trạng thái để kết quả có thể lặp lại.')
P('Defect được quản lý theo vòng đời New, Assigned, In progress, Fixed, Re-test, Closed hoặc Deferred. QA chỉ đóng defect khi bản sửa được xác nhận và regression liên quan không phát sinh lỗi mới. Defect Deferred phải có lý do, người chấp nhận, ảnh hưởng, workaround và thời điểm xem xét lại. Cách này giúp Test Summary phản ánh đúng rủi ro còn lại thay vì chỉ đếm số case pass.')
H('4.5. Thiết lập danh mục kiểm tra',2)
P('Mười checklist dưới đây được thiết kế theo đúng phạm vi của StayEasy. Mỗi checklist có người thực hiện, kết quả và bằng chứng; trạng thái được ghi theo các giá trị Đạt, Không đạt hoặc Không áp dụng.')
P('Checklist được thực hiện ở thời điểm phù hợp với Waterfall: SRS trước khi chốt yêu cầu; CSDL và coding trước baseline thiết kế/lập trình; test case trước khi thực thi; UI/UX và bảo mật trong kiểm thử; triển khai trước release; đánh giá cuối và QA test trước acceptance; code review trước khi merge. Người thực hiện ghi kết quả, ngày, version và liên kết bằng chứng trong hồ sơ dự án.')
H('4.5.1. Checklist review tài liệu yêu cầu (SRS)',3)
tbl('Bảng 8. Checklist review SRS',['TT','Nội dung kiểm tra','Kết quả','Bằng chứng'], [['1','Mỗi yêu cầu có mã duy nhất và mô tả có thể kiểm thử.','...','SRS/RTM'],['2','BR, SR, NFR không trùng hoặc mâu thuẫn.','...','SRS review'],['3','Nhóm người dùng và quyền được xác định rõ.','...','SRS'],['4','Tiêu chí chấp nhận có thể quan sát hoặc đo được.','...','Acceptance criteria'],['5','Yêu cầu có liên kết tới module hoặc test case.','...','RTM'],['6','Thuật ngữ StayEasy được dùng thống nhất.','...','Review record']],[0.45,3.7,0.8,1.7])
H('4.5.2. Checklist thiết kế CSDL',3)
tbl('Bảng 9. Checklist thiết kế CSDL',['TT','Nội dung kiểm tra','Kết quả','Bằng chứng'], [['1','Các bảng KHACHHANG, QUANTRIVIEN, LOAIPHONG, CANPHONG được mô tả rõ.','...','ERD/database.sql'],['2','DATPHONG có liên kết đúng khách, căn, thời gian và trạng thái.','...','ERD'],['3','HOSOLUUTRU và LICHSULUUTRU có quan hệ và owner rõ.','...','ERD'],['4','Có PK, FK, unique/index và kiểu dữ liệu phù hợp.','...','DDL/review'],['5','Ràng buộc không tạo dữ liệu mồ côi sau update/delete.','...','DB test'],['6','Có dữ liệu seed và hướng dẫn import/backup.','...','README']],[0.45,3.7,0.8,1.7])
H('4.5.3. Checklist coding standard backend',3)
tbl('Bảng 10. Checklist coding standard backend',['TT','Tiêu chí','Kết quả','Bằng chứng'], [['1','Tên biến, hàm, route và module theo quy ước.','...','Code review'],['2','Input được validation trước xử lý.','...','server.js/test'],['3','Truy vấn MySQL sử dụng parameterized query.','...','Code review'],['4','I/O, database và API có xử lý lỗi.','...','Code/test'],['5','Không log mật khẩu, token hoặc secret.','...','Log review'],['6','Middleware auth/admin/ownership được dùng đúng route.','...','Security test'],['7','Thay đổi có unit/integration test liên quan.','...','Test record']],[0.45,3.7,0.8,1.7])
H('4.5.4. Checklist test case chức năng đặt phòng',3)
tbl('Bảng 11. Checklist test case chức năng đặt phòng',['TT','Tiêu chí','Kết quả','Bằng chứng'], [['1','Có case đặt phòng hợp lệ với ngày và căn hợp lệ.','...','TC BOOK'],['2','Có case thiếu trường bắt buộc hoặc sai định dạng.','...','TC BOOK'],['3','Có case trùng căn và khoảng thời gian.','...','TC BOOK'],['4','Có case khách xem đơn của khách khác.','...','Security test'],['5','Có expected cho HTTP response, giao diện và DB.','...','Execution log'],['6','Có actual result, pass/fail và screenshot/log.','...','Evidence']],[0.45,3.7,0.8,1.7])
H('4.5.5. Checklist UI/UX',3)
tbl('Bảng 12. Checklist UI/UX',['TT','Tiêu chí','Kết quả','Bằng chứng'], [['1','Luồng đăng nhập và đặt phòng có điều hướng rõ.','...','Screenshot'],['2','Form có validation và thông báo lỗi dễ hiểu.','...','UI test'],['3','Không tràn hoặc che khuất ở desktop/mobile viewport.','...','Cross-browser'],['4','Trạng thái loading, empty, error và success được thể hiện.','...','UI test'],['5','Màu, chữ, nút và khoảng cách nhất quán.','...','UI review'],['6','Không hiển thị dữ liệu nhạy cảm không cần thiết.','...','Security UI']],[0.45,3.7,0.8,1.7])
H('4.5.6. Checklist bảo mật (QE)',3)
tbl('Bảng 13. Checklist bảo mật (QE)',['TT','Tiêu chí','Kết quả','Bằng chứng'], [['1','Token sai/thiếu không truy cập route bảo vệ.','...','Security test'],['2','Customer không gọi được API admin.','...','Negative test'],['3','Customer A không xem/sửa dữ liệu customer B.','...','IDOR test'],['4','Mật khẩu được hash có salt và không ghi log.','...','Code/log review'],['5','Input và query chống SQLi/XSS cơ bản.','...','OWASP ZAP/checklist'],['6','Rate limit đăng nhập và CORS được kiểm tra.','...','Security log']],[0.45,3.7,0.8,1.7])
H('4.5.7. Checklist triển khai',3)
tbl('Bảng 14. Checklist triển khai',['TT','Tiêu chí','Kết quả','Bằng chứng'], [['1','Node.js, Python/FastAPI và MySQL đúng phiên bản.','...','Environment record'],['2','Biến môi trường và secret được cấu hình ngoài repository.','...','Config checklist'],['3','database.sql import thành công và có dữ liệu seed.','...','Import log'],['4','Backend, frontend và chatbot khởi động đúng cổng.','...','Smoke test'],['5','Có backup trước release và hướng dẫn restore.','...','Backup record'],['6','README, package, requirements và release note đầy đủ.','...','Handover checklist']],[0.45,3.7,0.8,1.7])
H('4.5.8. Checklist đánh giá chất lượng cuối dự án',3)
tbl('Bảng 15. Checklist đánh giá chất lượng cuối dự án',['TT','Tiêu chí','Kết quả','Bằng chứng'], [['1','BR/SR/NFR chính có test case liên kết.','...','RTM'],['2','Không còn blocker/critical mở trước acceptance.','...','Defect Report'],['3','Code review và tài liệu baseline đã được phê duyệt.','...','Review record'],['4','Security, compatibility và usability đã được kiểm tra.','...','Test Summary'],['5','UAT có kết quả và chữ ký xác nhận.','...','UAT record'],['6','Residual risk có owner và kế hoạch xử lý.','...','Risk log']],[0.45,3.7,0.8,1.7])
H('4.5.9. QA test checklist',3)
tbl('Bảng 16. QA test checklist',['TT','Tiêu chí','Kết quả','Bằng chứng'], [['1','Test Plan xác định phạm vi, môi trường và vai trò.','...','Test Plan'],['2','Test case có ID, tiền điều kiện, data, expected và yêu cầu liên kết.','...','Test Cases'],['3','Unit, integration, system, security và acceptance được bao phủ.','...','Test matrix'],['4','Regression được thực hiện sau sửa defect.','...','Re-test log'],['5','Actual result và evidence được lưu.','...','Execution log'],['6','Test Summary ghi pass/fail, defect và rủi ro tồn đọng.','...','Test Summary']],[0.45,3.7,0.8,1.7])
H('4.5.10. Code review checklist',3)
tbl('Bảng 17. Code review checklist',['TT','Tiêu chí','Ưu tiên','Kết quả'], [['1','Tên biến/hàm rõ nghĩa, đúng convention.','Cao','...'],['2','Không lặp code, code chết hoặc workaround không giải thích.','Trung bình','...'],['3','Validation, authorization và ownership đúng.','Cao','...'],['4','Parameterized query, không lộ lỗi hoặc secret.','Cao','...'],['5','Xử lý exception, status code và timeout phù hợp.','Cao','...'],['6','Có test tương ứng với thay đổi nghiệp vụ.','Cao','...'],['7','Diff nhỏ, có mô tả và reviewer độc lập.','Trung bình','...']],[0.45,3.7,0.8,1.7])

# 4.6 procedures
H('4.6. Thiết lập thủ tục QA',2)
H('4.6.1. Thủ tục kiểm tra mã nguồn (Code review procedures)',3)
P('Các bước thực hiện: (1) Developer xác định phạm vi thay đổi, branch, yêu cầu và test liên quan; (2) Developer chạy self-check, unit test và kiểm tra secret; (3) reviewer độc lập đọc diff, checklist và các điểm ảnh hưởng API/database/security; (4) ghi finding theo severity và owner; (5) developer sửa, bổ sung test và phản hồi; (6) reviewer/QA xác nhận, lưu review record và cho phép baseline.')
diagram('fig2_code_review','Hình 2. Sơ đồ luồng thủ tục kiểm tra mã nguồn', ['Self-check','Peer review','QA trace','Finding\n& fix','Re-review','Baseline'], cols=3)
P('Công cụ hỗ trợ: Git/GitHub, IDE, ESLint, SonarQube và Code_Review_Checklist. Vai trò: Developer tự kiểm tra và sửa; reviewer độc lập đánh giá; QA leader kiểm tra truy xuất và đóng review; SQA team audit bằng chứng; PM phê duyệt baseline khi thay đổi có ảnh hưởng lớn.')
P('Đầu vào của thủ tục là branch hoặc version xác định, yêu cầu liên quan, thiết kế liên quan và kết quả self-check. Đầu ra là review record có kết luận Approved, Approved with comments hoặc Rework required. Nếu finding liên quan tới bảo mật, quyền truy cập, dữ liệu hoặc API contract, reviewer phải yêu cầu kiểm thử bổ sung trước khi đóng. Review record được liên kết tới commit, file, dòng mã hoặc test case để người khác có thể xác minh.')
H('4.6.2. Thủ tục kiểm thử phần mềm (Software testing procedures)',3)
P('Các bước thực hiện: (1) QA chốt version, phạm vi và rủi ro; (2) chuẩn bị môi trường Node.js, MySQL, FastAPI, browser và tài khoản; (3) thiết kế hoặc cập nhật test case; (4) thực thi unit, integration, system, security, compatibility và acceptance; (5) ghi actual result/evidence; (6) mở defect khi expected và actual khác nhau; (7) re-test, regression, Test Summary và quyết định acceptance.')
diagram('fig3_testing','Hình 3. Sơ đồ luồng thủ tục kiểm thử phần mềm', ['Plan','Design\ntest case','Prepare\nenvironment','Execute','Defect\n& re-test','Summary\n/acceptance'], cols=3)
P('Công cụ hỗ trợ: Jest cho unit test, Postman/Newman cho API, JMeter cho hiệu năng, Playwright hoặc DevTools cho UI, OWASP ZAP cho kiểm tra bảo mật và Test_Plan/Test_Cases/Defect_Report để lưu bằng chứng. Vai trò: QA leader điều phối; Tester thực thi; Developer hỗ trợ unit và sửa lỗi; SQA team kiểm tra tính đầy đủ; khách hàng nội bộ thực hiện acceptance.')
P('Đầu vào của thủ tục là build đã được xác định và môi trường có thể lặp lại. Đầu ra tối thiểu gồm execution log, danh sách pass/fail/blocked, defect report, re-test log và Test Summary. Khi test bị blocked do môi trường hoặc dữ liệu, Tester phải ghi nguyên nhân và không tính blocked như pass. Khi kết quả khác expected, việc mở defect phải được thực hiện trước khi sửa trực tiếp để bảo đảm có lịch sử và bằng chứng.')
H('4.6.3. Thủ tục kiểm soát tài liệu (Document control procedures)',3)
P('Các bước thực hiện: (1) tác giả tạo tài liệu theo Doc_Stand_Guide; (2) gán mã, phiên bản, tác giả, reviewer và trạng thái; (3) reviewer kiểm tra nội dung, định dạng và liên kết; (4) QA/SQA ghi finding và theo dõi sửa; (5) PM hoặc người có thẩm quyền phê duyệt baseline; (6) lưu bản đã duyệt và bản cũ để truy xuất; (7) cập nhật change log khi tài liệu thay đổi.')
diagram('fig4_document','Hình 4. Sơ đồ luồng thủ tục kiểm soát tài liệu', ['Soạn thảo','Review\nnội dung','Sửa finding','Phê duyệt\nbaseline','Lưu trữ\n& truy xuất'], cols=3)
P('Công cụ hỗ trợ: Microsoft Word, Markdown, draw.io/Mermaid source, Git/GitHub và thư mục hồ sơ được phân quyền. Vai trò: tác giả chịu trách nhiệm nội dung; reviewer kiểm tra độc lập; QA leader theo dõi trạng thái; SQA team audit version history; PM phê duyệt baseline.')
P('Một tài liệu chỉ được xem là kiểm soát tốt khi người đọc xác định được bản nào đang có hiệu lực, bản nào đã thay thế, ai phê duyệt và những tài liệu nào bị ảnh hưởng. Doc_Stand_Guide quy định cách đặt tên, thư mục, số phiên bản và cách lưu source diagram. Khi phát hiện tài liệu dùng sai version, SQA team ghi finding, xác định phạm vi ảnh hưởng và yêu cầu phát hành lại bằng chứng liên quan.')
H('4.6.4. Thủ tục kiểm soát thay đổi (Change control procedures)',3)
P('Các bước thực hiện: (1) ghi nhận Change Request; (2) mô tả lý do, phạm vi, ưu tiên và người yêu cầu; (3) BA, trưởng nhóm thiết kế, Developer và QA phân tích ảnh hưởng tới SRS, SDD, API, CSDL, security và regression; (4) PM/khách hàng nội bộ phê duyệt hoặc từ chối; (5) triển khai trên branch, review và kiểm thử; (6) cập nhật tài liệu, RTM, test case, release note và đóng CR.')
diagram('fig5_change','Hình 5. Sơ đồ luồng thủ tục kiểm soát thay đổi', ['Change\nrequest','Impact\nanalysis','Approve\nor reject','Implement\nbranch','Review &\nregression','Close CR\n& baseline'], cols=3)
P('Công cụ hỗ trợ: Git/GitHub Issues, pull request, RTM, change log, review checklist và Test Plan. Vai trò: người yêu cầu mô tả thay đổi; BA phân tích yêu cầu; trưởng nhóm thiết kế phân tích kiến trúc/dữ liệu; Developer thực hiện; QA kiểm thử hồi quy; PM/khách hàng nội bộ quyết định.')
P('Điều kiện đóng Change Request là quyết định được phê duyệt, code và tài liệu đã cập nhật, review đã hoàn thành, test case liên quan đã chạy, defect phát sinh đã có trạng thái và release note đã ghi thay đổi. Nếu thay đổi bị từ chối, CR vẫn được lưu với lý do. Nếu thay đổi được chấp nhận có điều kiện, điều kiện, người chấp nhận và hạn xem xét phải được lưu cùng CR.')
H('4.6.5. Thủ tục kiểm thử mã nguồn tự động (Automated code testing procedures)',3)
P('Các bước thực hiện: (1) xác định logic cần tự động hóa; (2) viết unit/integration test có dữ liệu độc lập; (3) chạy ESLint hoặc kiểm tra syntax; (4) chạy Jest và kiểm tra kết quả; (5) lưu log và artifact; (6) pipeline thất bại thì tạo finding/defect; (7) chỉ cho phép merge khi điều kiện chất lượng được đáp ứng hoặc có ngoại lệ được phê duyệt.')
diagram('fig6_automation','Hình 6. Sơ đồ luồng thủ tục kiểm thử mã nguồn tự động', ['Commit/PR','Install\ndependency','Lint &\nsyntax','Jest\nunit/integration','Publish\nlog/artifact','Gate\npass/fail'], cols=3)
P('Công cụ hỗ trợ: ESLint, Jest, npm scripts, Jenkins hoặc GitHub Actions, SonarQube và artifact storage. Vai trò: Developer viết test và xử lý lỗi; QA leader định nghĩa gate; SQA team kiểm tra log; PM phê duyệt ngoại lệ; trưởng nhóm thiết kế xem xét ảnh hưởng kiến trúc.')
P('Bộ kiểm thử tự động nên ưu tiên các logic có rủi ro hồi quy cao: hash/verify mật khẩu, xác thực JWT, kiểm tra role và ownership, validate ngày, kiểm tra điều kiện đặt phòng, xử lý response lỗi và các helper dùng chung. Test phải độc lập với dữ liệu production, có seed hoặc mock rõ ràng, không phụ thuộc thứ tự chạy. Log pipeline phải cho biết commit, runtime, lệnh chạy, kết quả và artifact để QA có thể kiểm tra lại.')
H('4.6.6. Thủ tục báo cáo và phản hồi (Reporting and feedback procedures)',3)
P('Các bước thực hiện: (1) QA thu thập tiến độ review/test, defect và rủi ro; (2) chuẩn hóa báo cáo theo template; (3) reviewer kiểm tra số liệu và bằng chứng; (4) gửi báo cáo cho PM, SQA team và các vai trò liên quan; (5) tổ chức phản hồi, xác định hành động; (6) cập nhật trạng thái, owner và thời hạn; (7) lưu báo cáo vào baseline hoặc hồ sơ dự án.')
diagram('fig7_reporting','Hình 7. Sơ đồ luồng thủ tục báo cáo và phản hồi', ['Collect\ndata','Validate\nevidence','Report\nQA status','Review &\nfeedback','Action\ntracking'], cols=3)
P('Công cụ hỗ trợ: Test_Summary_Report, Defect_Report, dashboard, GitHub Issues, email hoặc biên bản họp được lưu trong hồ sơ. Vai trò: QA leader lập báo cáo; Tester/Developer cung cấp bằng chứng; SQA team kiểm tra tính nhất quán; PM quyết định hành động; khách hàng nội bộ phản hồi về nghiệp vụ và acceptance.')
P('Báo cáo QA phải phân biệt sự kiện, kết quả đo, rủi ro và quyết định. Không dùng các từ như “đã ổn” nếu không có evidence hoặc version. Báo cáo tiến độ ghi phạm vi đã thực thi và phạm vi còn lại; Defect Report mô tả sai khác; Test Summary kết luận theo tiêu chí; phản hồi người dùng ghi rõ yêu cầu hoặc luồng bị ảnh hưởng. Cách trình bày này giúp PM và khách hàng nội bộ ra quyết định dựa trên thông tin có thể kiểm tra.')
H('4.7. Thiết lập quy trình QA',2)
P('Quy trình QA kết hợp review tĩnh và kiểm thử động. Review tĩnh nhằm phát hiện lỗi trong SRS, SDD, code, checklist và tài liệu trước khi chạy hệ thống. Kiểm thử động nhằm xác nhận hành vi thực tế của API, database, giao diện, phân quyền, hiệu năng và luồng nghiệp vụ. Mục tiêu là phát hiện sớm, truy xuất được, ngăn lỗi tái diễn và cung cấp bằng chứng cho quyết định release.')
P('Các chỉ số theo dõi quy trình QA gồm số lượng yêu cầu có test liên kết, tỷ lệ review đã đóng finding, số case đã thực thi, tỷ lệ pass/fail, số defect theo severity, thời gian từ mở tới đóng defect, số defect tái mở, phạm vi regression, tình trạng checklist và số hành động cải tiến quá hạn. Chỉ số phải được hiểu trong bối cảnh phạm vi và version; không dùng tỷ lệ pass đơn lẻ để kết luận chất lượng nếu còn blocker/critical hoặc thiếu evidence.')
P('Quy trình QA cũng phải kiểm soát giao tiếp giữa các vai trò. Developer cung cấp code, log và giải thích kỹ thuật; QA cung cấp test result, defect và rủi ro; SQA team kiểm tra tuân thủ; PM quyết định ưu tiên và baseline; khách hàng nội bộ xác nhận nghiệp vụ. Các quyết định quan trọng phải được lưu thành biên bản, issue, review record hoặc sign-off để tránh phụ thuộc vào trao đổi miệng.')
P('Ở cổng yêu cầu, QA kiểm tra sự đầy đủ và khả kiểm thử của SRS. Ở cổng thiết kế, QA đối chiếu SDD, ERD, API contract và các yêu cầu bảo mật. Ở cổng code, QA kiểm tra review record, unit test và thay đổi so với baseline. Ở cổng kiểm thử, QA kiểm tra coverage, defect lifecycle, regression và kết quả đo. Ở cổng release, QA kiểm tra UAT, residual risk, hướng dẫn triển khai, backup/restore và quyết định phê duyệt. Các cổng này tạo thành chuỗi bằng chứng liên tục.')
P('Khi một quy trình phát hiện lỗi, thông tin được chuyển về đúng nơi: lỗi tài liệu tạo finding; lỗi mã nguồn tạo defect; yêu cầu mới tạo Change Request; thiếu bằng chứng tạo action item; phản hồi nghiệp vụ tạo cập nhật acceptance criteria. Việc phân loại đúng giúp nhóm không sửa sai artifact hoặc bỏ qua tác động tới test và dữ liệu.')
P('Trong mỗi milestone, SQA team thực hiện một lần kiểm tra chéo giữa yêu cầu, thiết kế, mã nguồn, test case và bằng chứng. Kiểm tra chéo trả lời bốn câu hỏi: yêu cầu đã có cách hiện thực chưa; cách hiện thực đã có review chưa; test case đã kiểm tra đúng rủi ro chưa; và kết quả có thể truy xuất về version nào. Nếu một câu trả lời là chưa, trạng thái cổng chất lượng được ghi là cần bổ sung và hành động khắc phục được giao cho vai trò phù hợp.')
P('Hồ sơ QA hoàn chỉnh của StayEasy gồm SRS/RTM, SDD/ERD/API contract, review record, Test Plan, Test Cases, execution log, Defect Report, re-test/regression evidence, Test Summary, UAT record, release note và Retrospective. Các hồ sơ này không chỉ phục vụ việc chấm báo cáo mà còn tạo cơ sở để bảo trì, điều tra lỗi sau bàn giao và đánh giá tác động khi hệ thống mở rộng thêm thanh toán, nhiều chi nhánh hoặc tích hợp OTA trong tương lai.')
P('SQA audit được tổ chức theo mốc: trước baseline yêu cầu, trước baseline thiết kế, trước merge code, trước kiểm thử acceptance và trước release. Mỗi audit lấy mẫu một số yêu cầu để truy theo chiều dọc từ SRS tới SDD, module, test case, kết quả và defect; đồng thời lấy mẫu một số thay đổi để truy theo chiều ngang từ CR tới review, regression và release note. Kết quả audit được phân loại thành phù hợp, cần khắc phục hoặc không phù hợp và được theo dõi tới khi đóng.')
P('Cải tiến quy trình được ghi nhận trong Retrospective bằng cấu trúc hiện trạng, nguyên nhân, tác động, hành động, owner và thời hạn. Các bài học có thể bao gồm bổ sung test cho điều kiện biên, chuẩn hóa response API, cải thiện dữ liệu seed, tự động hóa smoke test hoặc cập nhật checklist bảo mật. Chỉ hành động có owner và thời hạn mới được xem là hành động cải tiến có thể theo dõi.')
P('Khi đánh giá kết quả QA, nhóm cần xem đồng thời chất lượng sản phẩm và chất lượng quy trình. Một sản phẩm có tỷ lệ test pass cao nhưng thiếu review, thiếu dữ liệu âm, không có regression hoặc còn lỗ hổng quyền truy cập vẫn chưa đạt yêu cầu chất lượng. Ngược lại, một defect được ghi nhận rõ, có severity, owner, workaround và quyết định chấp nhận có điều kiện sẽ giúp rủi ro được quản lý minh bạch hơn so với việc bỏ qua defect để làm đẹp số liệu.')
P('Các cuộc họp QA nên tập trung vào quyết định và hành động: phạm vi đã hoàn thành, rủi ro mới, defect cần ưu tiên, thay đổi ảnh hưởng baseline, bằng chứng còn thiếu và quyết định cần PM hoặc khách hàng nội bộ phê duyệt. Biên bản ghi rõ người phát biểu, kết luận, owner và thời hạn. Đây là phần quan trọng của SQA vì bảo đảm thông tin chất lượng đi tới đúng người có quyền quyết định.')
P('Năm quy trình QA cần xây dựng cho StayEasy gồm: (1) quy trình review yêu cầu/thiết kế/code; (2) quy trình kiểm thử và quản lý defect; (3) quy trình kiểm soát tài liệu và baseline; (4) quy trình kiểm soát thay đổi và regression; (5) quy trình báo cáo, phản hồi và cải tiến. Các quy trình liên kết theo Waterfall nhưng có thể tạo vòng phản hồi có kiểm soát khi phát hiện defect hoặc thay đổi.')
diagram('fig8_qa_relationship','Hình 8. Mối quan hệ giữa các quy trình QA', ['SRS/SDD\nBaseline','Review\ntĩnh','Code &\nTest','Defect/CR\ncontrol','Regression &\nsummary','Release/UAT\nfeedback'], cols=3)
H('4.8. Thiết lập danh mục kiểm tra',2)
P('Checklist giúp biến yêu cầu chất lượng thành các điểm kiểm tra có thể lặp lại, giảm bỏ sót, phân định trách nhiệm, tạo bằng chứng và hỗ trợ audit. Checklist phải ngắn gọn, có phiên bản, người thực hiện, thời điểm, kết quả và liên kết tới artifact liên quan; không thay thế cho Test Case hoặc tài liệu thiết kế.')
P('Khi sử dụng checklist, người thực hiện đọc từng mục theo thứ tự, đánh dấu Đạt/Không đạt/Không áp dụng, ghi chú bằng chứng và mở finding nếu không đạt. QA leader kiểm tra tính hợp lệ của bằng chứng; SQA team kiểm tra việc thực hiện đúng phiên bản checklist; PM hoặc khách hàng nội bộ xác nhận những mục thuộc quyền quyết định của mình. Checklist được cập nhật khi quy trình, công nghệ hoặc rủi ro của dự án thay đổi.')
P('Checklist test được rà soát trước mỗi đợt test; checklist code review được dùng cho mọi pull request có thay đổi chính; checklist tài liệu được dùng khi tạo mới, sửa hoặc đưa artifact vào baseline. Những mục Không áp dụng phải có lý do, vì một mục bị bỏ qua mà không giải thích có thể tạo lỗ hổng trong audit. Sau mỗi milestone, SQA team tổng hợp các mục thường xuyên Không đạt để chuyển thành hành động cải tiến.')
bullets(['Test checklist: tập trung vào phạm vi, môi trường, dữ liệu, cấp độ test, actual result, evidence, defect, re-test, regression và điều kiện kết thúc.','Code review checklist: tập trung vào readability, naming, cấu trúc, validation, exception, security, query, test, diff và khả năng bảo trì.','Checklist tài liệu: tập trung vào mã tài liệu, phiên bản, tác giả, reviewer, phê duyệt, tính nhất quán, truy xuất, thay đổi và lưu trữ.'])
page_break()

# 5
H('5. CÔNG CỤ HỖ TRỢ',1)
P('Các công cụ được phân nhóm theo hoạt động chất lượng. Việc lựa chọn công cụ phải phù hợp phạm vi StayEasy, khả năng lưu bằng chứng và quy trình Waterfall; công cụ không thay thế trách nhiệm review, test hoặc phê duyệt.')
P('Tiêu chí chọn công cụ gồm khả năng tích hợp với mã nguồn, khả năng chạy lặp lại, khả năng lưu log và artifact, phân quyền người dùng, chi phí, mức độ phù hợp với Node.js/MySQL/FastAPI và khả năng hỗ trợ audit. Mỗi công cụ cần có owner, phiên bản, hướng dẫn sử dụng và nơi lưu kết quả. Khi công cụ không khả dụng, nhóm ghi rõ phương án thay thế và không bỏ qua bước kiểm soát chất lượng.')
P('Việc kết hợp công cụ được tổ chức theo luồng: Git/GitHub quản lý source và thay đổi; SonarQube/ESLint cung cấp tín hiệu chất lượng code; Jest và Postman/Newman cung cấp kết quả kiểm thử; JMeter và OWASP ZAP cung cấp evidence hiệu năng/bảo mật; Jenkins hoặc GitHub Actions tự động hóa việc chạy và lưu artifact; Word/Markdown lưu tài liệu và báo cáo. Mỗi kết quả phải liên kết với version hoặc commit để tránh trộn dữ liệu giữa các lần chạy.')
P('SonarQube và ESLint được sử dụng ở thời điểm review mã nguồn để phát hiện lỗi quy ước, code smell, duplication và dấu hiệu rủi ro trước khi merge. Với backend, kết quả phải được đọc cùng với review thủ công vì công cụ tĩnh không thay thế việc kiểm tra ownership, trạng thái đơn hoặc tính đúng của nghiệp vụ. Với chatbot Python, kiểm tra style và dependency giúp giảm rủi ro nhưng vẫn cần kiểm thử timeout, response và fallback.')
P('Postman/Newman và Jest cung cấp hai lớp bằng chứng khác nhau. Postman/Newman kiểm tra hợp đồng API, status code, body, header và quyền truy cập theo vai trò; Jest kiểm tra helper, validation, middleware và logic nhỏ với dữ liệu có thể lặp lại. JMeter được dùng khi cần đánh giá tải và p95, trong khi OWASP ZAP hỗ trợ phát hiện các vấn đề web phổ biến. Kết quả của các công cụ phải được lưu cùng build và môi trường.')
P('Jenkins hoặc GitHub Actions có thể tạo quality gate theo chuỗi: checkout commit, cài dependency, kiểm tra secret/config, chạy lint, chạy unit/integration test, tạo artifact và thông báo kết quả. Nếu pipeline fail, merge hoặc release phải dừng cho đến khi có bản sửa hoặc quyết định ngoại lệ. Ngoại lệ phải ghi lý do, phạm vi, người phê duyệt và thời hạn xem xét lại.')
bullets(['Theo dõi chất lượng mã nguồn: SonarQube để theo dõi code smell, bug, vulnerability, duplication và maintainability; ESLint để kiểm tra quy ước JavaScript; Pylint hoặc flake8 có thể dùng cho Python/FastAPI.','Quản lý tài liệu và cải tiến quy trình: Microsoft Word, Markdown, Git/GitHub, draw.io hoặc Mermaid source, issue/change log và thư mục baseline; Retrospective dùng để ghi bài học và hành động cải tiến.','Thu thập phản hồi người dùng: biểu mẫu UAT, biên bản họp, issue form, khảo sát nội bộ và Test_Summary_Report; phản hồi phải gắn với yêu cầu hoặc tiêu chí chấp nhận.','Quản lý kiểm thử: Postman/Newman cho REST API, Jest cho unit/integration, Playwright hoặc DevTools cho UI, JMeter cho hiệu năng, Test Plan và Test Cases cho điều phối thực thi.','Theo dõi lỗi: GitHub Issues hoặc công cụ defect tracker có mã lỗi, severity, priority, owner, trạng thái, bước tái hiện, evidence, re-test và resolution.','Tự động hóa kiểm thử: Jenkins hoặc GitHub Actions để chạy lint, Jest, integration test, kiểm tra build và lưu artifact; OWASP ZAP hỗ trợ security scan; SonarQube có thể làm quality gate trước merge/release.'])
P('Với StayEasy, bộ tối thiểu gồm Git/GitHub, Word/Markdown, MySQL Workbench hoặc CLI, Postman/Newman, Jest, ESLint, JMeter, OWASP ZAP và Jenkins/GitHub Actions. Cấu hình chi tiết, phiên bản, dữ liệu và người chịu trách nhiệm phải được ghi trong QA_Plan và Test_Plan của dự án.')

d.core_properties.title = 'Báo cáo đảm bảo chất lượng phần mềm StayEasy Homestay theo khung Bài 08'
d.core_properties.subject = 'Waterfall, ISO/IEC 25010:2011 và hệ thống SQA'
d.core_properties.author = ''
d.save(OUT)
print(OUT)
